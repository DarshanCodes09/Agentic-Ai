"""
Tests for document processing: text cleaning, PDF/DOCX extraction, and chunking — Phase 3.
"""

import io
from pathlib import Path

import docx
import fitz
import pytest

from app.core.exceptions import AppException
from app.services.document_processing.chunker import chunk_text
from app.services.document_processing.docx_processor import extract_text_from_docx
from app.services.document_processing.pdf_processor import extract_text_from_pdf
from app.services.document_processing.text_cleaner import clean_text


# ---------------------------------------------------------------------------
# 1. Text Cleaning Tests
# ---------------------------------------------------------------------------

def test_clean_text_normalizes_whitespace_and_newlines():
    raw = "Line 1\r\n\r\n\r\n   Line 2 with   extra   spaces\t\tand tabs.   \n\n\n\nLine 3"
    cleaned = clean_text(raw)
    expected = "Line 1\n\nLine 2 with extra spaces and tabs.\n\nLine 3"
    assert cleaned == expected


def test_clean_text_strips_control_characters():
    raw = "Valid text\x00\x08with\x0bcontrol\x1fcharacters"
    cleaned = clean_text(raw)
    assert cleaned == "Valid textwithcontrolcharacters"


def test_clean_text_empty_input():
    assert clean_text("") == ""
    assert clean_text("   \n\n\t  ") == ""


# ---------------------------------------------------------------------------
# 2. PDF Extraction Tests (PyMuPDF)
# ---------------------------------------------------------------------------

def test_pdf_extraction_from_bytes():
    # Create an in-memory PDF with PyMuPDF
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), "Academic lecture on Distributed Systems and Consensus algorithms.")
    pdf_bytes = doc.tobytes()
    doc.close()

    text = extract_text_from_pdf(pdf_bytes)
    assert "Distributed Systems" in text
    assert "Consensus algorithms" in text


def test_pdf_extraction_from_file(tmp_path: Path):
    pdf_path = tmp_path / "sample.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), "Chapter 1: Relational Algebra and Normalization.")
    doc.save(str(pdf_path))
    doc.close()

    text = extract_text_from_pdf(pdf_path)
    assert "Relational Algebra" in text
    assert "Normalization" in text


def test_pdf_extraction_corrupt_file_raises():
    with pytest.raises(AppException):
        extract_text_from_pdf(b"Not a real PDF stream")


# ---------------------------------------------------------------------------
# 3. DOCX Extraction Tests (python-docx)
# ---------------------------------------------------------------------------

def test_docx_extraction_from_bytes():
    doc = docx.Document()
    doc.add_heading("Machine Learning Overview", level=1)
    doc.add_paragraph("Supervised learning utilizes labeled datasets to train models.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Algorithm"
    table.rows[0].cells[1].text = "Type"

    buffer = io.BytesIO()
    doc.save(buffer)
    docx_bytes = buffer.getvalue()

    text = extract_text_from_docx(docx_bytes)
    assert "Machine Learning Overview" in text
    assert "Supervised learning" in text
    assert "Algorithm | Type" in text


def test_docx_extraction_from_file(tmp_path: Path):
    docx_path = tmp_path / "notes.docx"
    doc = docx.Document()
    doc.add_paragraph("Graph theory: Dijkstra's algorithm finds shortest paths.")
    doc.save(str(docx_path))

    text = extract_text_from_docx(docx_path)
    assert "Dijkstra's algorithm" in text


def test_docx_extraction_corrupt_file_raises():
    with pytest.raises(AppException):
        extract_text_from_docx(b"Not a real DOCX file")


# ---------------------------------------------------------------------------
# 4. Deterministic Chunking Tests
# ---------------------------------------------------------------------------

def test_deterministic_chunking_with_metadata():
    text = (
        "Operating systems manage hardware resources. "
        "Memory virtualization provides processes with an illusion of dedicated memory. "
        "Concurrency introduces race conditions and requires locks, semaphores, or condition variables. "
        "Persistence stores data reliably using filesystems and journaling."
    )

    chunks1 = chunk_text(
        text=text,
        subject_id=10,
        course_material_id=42,
        source_filename="os_notes.pdf",
        chunk_size=100,
        chunk_overlap=20,
    )

    chunks2 = chunk_text(
        text=text,
        subject_id=10,
        course_material_id=42,
        source_filename="os_notes.pdf",
        chunk_size=100,
        chunk_overlap=20,
    )

    # Determinism: identical output across separate runs
    assert len(chunks1) > 1
    assert len(chunks1) == len(chunks2)
    for c1, c2 in zip(chunks1, chunks2):
        assert c1.text == c2.text
        assert c1.chunk_index == c2.chunk_index
        assert c1.subject_id == 10
        assert c1.course_material_id == 42
        assert c1.source_filename == "os_notes.pdf"
        assert c1.metadata == {
            "subject_id": 10,
            "course_material_id": 42,
            "source_filename": "os_notes.pdf",
            "chunk_index": c1.chunk_index,
        }


def test_chunking_empty_text():
    chunks = chunk_text(
        text="",
        subject_id=1,
        course_material_id=1,
        source_filename="empty.pdf",
    )
    assert chunks == []
