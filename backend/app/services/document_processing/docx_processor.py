"""
DOCX text extractor using python-docx.
"""

import io
from pathlib import Path

from docx import Document

from app.core.exceptions import AppException
from app.services.document_processing.text_cleaner import clean_text


def extract_text_from_docx(file_source: str | Path | bytes) -> str:
    """
    Extract text content from a DOCX file path or bytes using python-docx.

    Args:
        file_source: File path or raw DOCX bytes.

    Returns:
        Cleaned extracted text string.

    Raises:
        AppException: If file cannot be read or parsed as a valid DOCX document.
    """
    try:
        if isinstance(file_source, bytes):
            stream = io.BytesIO(file_source)
            doc = Document(stream)
        else:
            path = Path(file_source)
            if not path.exists():
                raise AppException(f"DOCX file not found at: {path}", status_code=404)
            doc = Document(str(path))

        elements: list[str] = []

        # Extract text from paragraphs
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if text:
                elements.append(text)

        # Extract text from tables
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    elements.append(" | ".join(row_cells))

        full_raw_text = "\n\n".join(elements)
        return clean_text(full_raw_text)

    except Exception as exc:
        if isinstance(exc, AppException):
            raise
        raise AppException(f"Failed to extract text from DOCX: {exc}", status_code=422) from exc
